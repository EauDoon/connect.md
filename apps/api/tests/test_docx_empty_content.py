from __future__ import annotations

import io
import socket
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from fastapi import HTTPException
from PIL import Image

from app.docx_text import _join, docx_body_has_content, extract_docx_text
from app.ingest import _convert_binary


def save(document, tmp_path):
    source = tmp_path / "source.docx"
    document.save(source)
    return source


def force_fallback(monkeypatch):
    import markitdown

    monkeypatch.setattr(
        markitdown.MarkItDown, "convert_local", lambda *_: SimpleNamespace(text_content="")
    )


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def reject(*args, **kwargs):
        pytest.fail("DOCX conversion attempted network access")

    monkeypatch.setattr(socket.socket, "connect", reject)
    monkeypatch.setattr(socket, "create_connection", reject)


@pytest.mark.parametrize("columns", [1, 3, 50])
@pytest.mark.parametrize("whitespace", ["", " \t\n ", "\u00a0\u2003\u2028"])
@pytest.mark.parametrize("fallback", [False, True])
def test_blank_tables_fail_closed_in_actual_primary_and_fallback(
    tmp_path, monkeypatch, columns, whitespace, fallback
):
    document = Document()
    table = document.add_table(rows=2, cols=columns)
    for row in table.rows:
        for cell in row.cells:
            cell.text = whitespace
    document.add_paragraph(whitespace)
    source = save(document, tmp_path)
    contents = source.read_bytes()
    if fallback:
        force_fallback(monkeypatch)
    assert not docx_body_has_content(source)
    assert extract_docx_text(source, 0) == ""
    for maximum in (0, 10000):
        with pytest.raises(HTTPException) as failure:
            _convert_binary(contents, ".docx", max_extracted_bytes=maximum)
        assert failure.value.status_code == 422
        assert failure.value.detail["provenance"]["converter"] == "none"
    assert source.read_bytes() == contents


@pytest.mark.parametrize("column,expected", [(0, "é |  |"), (1, "| é |"), (2, "|  | é")])
def test_mixed_rows_preserve_empty_columns_at_exact_normalized_limits(
    tmp_path, monkeypatch, column, expected
):
    force_fallback(monkeypatch)
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(1, column).text = "  é  "
    source = save(document, tmp_path)
    maximum = len(expected.encode())
    assert extract_docx_text(source, maximum) == expected
    assert _convert_binary(source.read_bytes(), ".docx", max_extracted_bytes=maximum)[:2] == (
        expected, "python-docx-local"
    )
    with pytest.raises(ValueError, match="limit"):
        extract_docx_text(source, maximum - 1)
    with pytest.raises(HTTPException):
        _convert_binary(source.read_bytes(), ".docx", max_extracted_bytes=maximum - 1)


def test_blank_nested_merged_rows_and_paragraphs_preserve_order(tmp_path):
    document = Document()
    document.add_paragraph("Before")
    blank = document.add_table(rows=3, cols=3)
    blank.cell(0, 0).merge(blank.cell(2, 2)).add_table(rows=3, cols=3)
    table = document.add_table(rows=2, cols=3)
    cell = table.cell(0, 1)
    cell.text = "Top"
    nested = cell.add_table(rows=2, cols=3)
    nested.cell(0, 0).add_table(rows=2, cols=3)
    cell.add_paragraph(" \t ")
    cell.add_paragraph("Bottom")
    table.cell(1, 1).add_table(rows=2, cols=3).cell(1, 1).text = "Nested"
    document.add_paragraph("After")
    source = save(document, tmp_path)
    expected = "Before\n| Top Bottom |\n| | Nested | |\nAfter"
    assert extract_docx_text(source, len(expected.encode())) == expected


@pytest.mark.parametrize("name", [
    "header", "footer", "first_page_header", "first_page_footer",
    "even_page_header", "even_page_footer",
])
@pytest.mark.parametrize("has_text", [False, True])
def test_actual_primary_scaffolding_cannot_hide_active_header_footer_text(tmp_path, name, has_text):
    document = Document()
    document.add_table(rows=2, cols=3)
    section = document.sections[0]
    section.different_first_page_header_footer = True
    document.settings.odd_and_even_pages_header_footer = True
    variant = getattr(section, name)
    variant.add_table(rows=2, cols=3, width=Inches(4))
    if has_text:
        variant.paragraphs[0].text = "Sentinel"
    source = save(document, tmp_path)
    if has_text:
        assert _convert_binary(source.read_bytes(), ".docx", max_extracted_bytes=8)[:2] == (
            "Sentinel", "python-docx-local"
        )
        with pytest.raises(HTTPException):
            _convert_binary(source.read_bytes(), ".docx", max_extracted_bytes=7)
    else:
        assert extract_docx_text(source, 0) == ""
        with pytest.raises(HTTPException):
            _convert_binary(source.read_bytes(), ".docx")


def test_inactive_header_does_not_admit_empty_primary_table(tmp_path):
    document = Document()
    document.add_table(rows=2, cols=3)
    document.sections[0].first_page_header.paragraphs[0].text = "Inactive"
    source = save(document, tmp_path)
    with pytest.raises(HTTPException):
        _convert_binary(source.read_bytes(), ".docx")


@pytest.mark.parametrize("content", ["|", "---", "| --- |", "é"])
def test_actual_primary_retains_literal_punctuation_and_unicode(tmp_path, content):
    document = Document()
    document.add_table(rows=1, cols=3).cell(0, 1).text = content
    source = save(document, tmp_path)
    text, converter, warnings = _convert_binary(source.read_bytes(), ".docx")
    assert content in text
    assert converter == "markitdown-local"
    assert warnings == []


@pytest.mark.parametrize("kind", ["hyperlink", "image", "math"])
def test_actual_primary_retains_content_outside_plain_paragraph_runs(tmp_path, monkeypatch, kind):
    import markitdown

    document = Document()
    paragraph = document.add_table(rows=1, cols=2).cell(0, 0).paragraphs[0]
    if kind == "image":
        picture = io.BytesIO()
        Image.new("RGB", (4, 4), "white").save(picture, format="PNG")
        picture.seek(0)
        paragraph.add_run().add_picture(picture)
        expected = "!["
    elif kind == "math":
        math = OxmlElement("m:oMath")
        run = OxmlElement("m:r")
        text = OxmlElement("m:t")
        text.text = "x"
        run.append(text)
        math.append(run)
        paragraph._p.append(math)
        expected = "x"
    else:
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), "destination")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Link sentinel"
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        expected = "Link sentinel"
    source = save(document, tmp_path)
    assert docx_body_has_content(source)
    primary_text = markitdown.MarkItDown().convert_local(str(source)).text_content
    text, converter, _ = _convert_binary(source.read_bytes(), ".docx")
    assert converter == "markitdown-local"
    assert expected in text
    assert text == primary_text
    if kind == "image":
        force_fallback(monkeypatch)
        assert extract_docx_text(source, 0) == ""
        with pytest.raises(HTTPException):
            _convert_binary(source.read_bytes(), ".docx")


def test_join_preserves_internal_whitespace_but_bounds_final_text():
    assert _join(iter(["  A  ", " B  "]), " ", 6) == "A    B"
    assert _join(iter(["", "é", ""]), " | ", 6) == "| é |"
    assert _join(iter(["A", " " * 10000]), " ", 1) == "A"
    with pytest.raises(ValueError, match="limit"):
        _join(iter(["A", " " * 10000, "B"]), " ", 100)
