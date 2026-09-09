from __future__ import annotations

import io
import socket
import zipfile
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi import HTTPException

from app.config import Settings
from app.docx_text import extract_docx_text
from app.ingest import _convert_binary, _validate_file_signature


def document_bytes():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Ada Lovelace"
    document.add_paragraph("First paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "Systems"
    document.add_paragraph("Last paragraph")
    document.sections[0].footer.paragraphs[0].text = "Portfolio"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_actual_docx_fallback_is_offline_and_preserves_order(monkeypatch):
    import markitdown

    def unavailable(*args, **kwargs):
        raise RuntimeError("primary unavailable")

    def reject_network(*args, **kwargs):
        pytest.fail("local DOCX fallback attempted network access")

    monkeypatch.setattr(markitdown.MarkItDown, "convert_local", unavailable)
    monkeypatch.setattr(socket.socket, "connect", reject_network)
    monkeypatch.setattr(socket, "create_connection", reject_network)
    contents = document_bytes()
    _validate_file_signature(
        contents, ".docx", Settings(api_key_pepper="test-only-pepper-is-long-enough")
    )
    text, converter, warnings = _convert_binary(contents, ".docx")
    assert text == "Ada Lovelace\nFirst paragraph\nPython | Systems\nLast paragraph\nPortfolio"
    assert converter == "python-docx-local"
    assert warnings == ["MarkItDown conversion was unavailable: RuntimeError."]


def test_docx_fallback_rejects_utf8_overflow_without_a_partial_draft(monkeypatch):
    import markitdown

    monkeypatch.setattr(
        markitdown.MarkItDown, "convert_local", lambda *_: SimpleNamespace(text_content="")
    )
    with pytest.raises(HTTPException) as failure:
        _convert_binary(document_bytes(), ".docx", max_extracted_bytes=12)
    assert failure.value.status_code == 422
    assert failure.value.detail["provenance"]["converter"] == "none"


def test_docx_exact_utf8_limit_and_nested_merged_tables(tmp_path):
    document = Document()
    document.add_paragraph("é")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested"
    source = tmp_path / "source.docx"
    document.save(source)
    expected = "é\nMerged Nested"
    text = extract_docx_text(source, len(expected.encode()))
    assert text == expected
    with pytest.raises(ValueError, match="limit"):
        extract_docx_text(source, len(expected.encode()) - 1)


def test_docx_xml_does_not_resolve_external_entities(tmp_path):
    sentinel = tmp_path / "sentinel.txt"
    sentinel.write_text("SYNTHETIC_ENTITY_SENTINEL", encoding="utf-8")
    source = tmp_path / "source.docx"
    xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<!DOCTYPE w:document [<!ENTITY external SYSTEM "{sentinel.as_uri()}">]>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:body><w:p><w:r><w:t>&external;</w:t></w:r></w:p></w:body></w:document>'
    )
    with zipfile.ZipFile(io.BytesIO(document_bytes())) as original:
        with zipfile.ZipFile(source, "w") as output:
            for member in original.infolist():
                output.writestr(
                    member,
                    xml.encode() if member.filename == "word/document.xml" else original.read(member),
                )
    assert "SYNTHETIC_ENTITY_SENTINEL" not in extract_docx_text(source, 1024)
